import os
import io
from fastapi import FastAPI, UploadFile, File, BackgroundTasks
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import StreamingResponse
from typing import List
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage

try:
    from backend.pdf_processor import extract_text_from_pdf, IMAGE_DIR
    from backend.analyzer import clean_ocr_lines, extract_subject_info, split_into_questions_and_crop, group_questions
except ImportError:
    from pdf_processor import extract_text_from_pdf, IMAGE_DIR
    from analyzer import clean_ocr_lines, extract_subject_info, split_into_questions_and_crop, group_questions

app = FastAPI(title="PYQ Analyzer API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Serve images directory
app.mount("/images", StaticFiles(directory=IMAGE_DIR), name="images")

# Global cache to store last analysis results for download
last_analysis_results = {}

@app.get("/health")
def health_check():
    return {"status": "ok"}

@app.post("/analyze")
async def analyze_papers(files: List[UploadFile] = File(...)):
    global last_analysis_results
    all_pages_raw = []
    total_text_pages = 0
    total_ocr_pages = 0
    questions_by_subject = {}

    for file in files:
        content = await file.read()
        result = extract_text_from_pdf(file_bytes=content, filename=file.filename)
        
        total_text_pages += result["text_pages"]
        total_ocr_pages += result["ocr_pages"]
        
        # Get first page lines to extract subject info
        first_page_lines = result["pages"][0]["lines"] if result["pages"] else []
        subject_name, course_code = extract_subject_info(first_page_lines, file.filename)
        subject_key = f"{subject_name} ({course_code})"
        
        if subject_key not in questions_by_subject:
            questions_by_subject[subject_key] = []
            
        for page in result["pages"]:
            # Combine raw text for the frontend
            raw_text = "\n".join([l["text"] for l in page["lines"]])
            all_pages_raw.append({
                "filename": page["filename"],
                "page_num": page["page_num"],
                "method": page["method"],
                "text": raw_text,
                "char_count": len(raw_text)
            })
            
            cleaned_lines = clean_ocr_lines(page["lines"])
            extracted_qs = split_into_questions_and_crop(
                cleaned_lines, 
                page["image_path"], 
                page["height"], 
                file.filename, 
                page["page_num"]
            )
            
            for q in extracted_qs:
                q["paper"] = file.filename
                q["page"] = page["page_num"]
                questions_by_subject[subject_key].append(q)

    total_pages = total_text_pages + total_ocr_pages
    summary_text = f"{total_pages} pages: {total_text_pages} text, {total_ocr_pages} OCR"

    analysis_results = {}
    total_questions = 0
    for subj, q_list in questions_by_subject.items():
        total_questions += len(q_list)
        grouped = group_questions(q_list)
        analysis_results[subj] = grouped

    last_analysis_results = analysis_results

    return {
        "summary": summary_text,
        "total_pages": total_pages,
        "text_pages": total_text_pages,
        "ocr_pages": total_ocr_pages,
        "paper_count": len(files),
        "total_questions_extracted": total_questions,
        "analysis": analysis_results,
        "pages": all_pages_raw
    }

@app.get("/export")
def export_excel(subject: str = ""):
    global last_analysis_results
    if not subject or subject not in last_analysis_results:
        return {"error": "Subject not found in analysis cache."}
        
    data = last_analysis_results[subject]
    wb = Workbook()
    
    # Remove default sheet
    wb.remove(wb.active)
    
    for category in ["Repeated", "Least repeated", "Remaining"]:
        items = data.get(category, [])
        if not items:
            continue
            
        # Sheet names can be max 31 chars
        ws = wb.create_sheet(title=category[:31])
        ws.append(["Question", "Count", "Marks", "Locations", "Image"])
        
        # Set column widths
        ws.column_dimensions['A'].width = 60
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 10
        ws.column_dimensions['D'].width = 30
        ws.column_dimensions['E'].width = 50
        
        for idx, q in enumerate(items, start=2):
            ws.cell(row=idx, column=1, value=q["text"])
            ws.cell(row=idx, column=2, value=q["count"])
            ws.cell(row=idx, column=3, value=q["marks"])
            ws.cell(row=idx, column=4, value=q["locations"])
            
            # Embed image
            if q.get("image_url"):
                img_path = "static" + q["image_url"]
                if os.path.exists(img_path):
                    try:
                        xl_img = OpenpyxlImage(img_path)
                        # Resize image to fit row
                        xl_img.width = min(xl_img.width, 300)
                        xl_img.height = min(xl_img.height, 100)
                        ws.add_image(xl_img, f"E{idx}")
                        ws.row_dimensions[idx].height = max(75, xl_img.height * 0.75)
                    except Exception:
                        ws.cell(row=idx, column=5, value="Image error")
    
    # If no sheets created (e.g. no data), create empty sheet
    if not wb.sheetnames:
        wb.create_sheet("Empty")
        
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    
    filename = f"{subject}_analysis.xlsx".replace(" ", "_")
    return StreamingResponse(
        out, 
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)
